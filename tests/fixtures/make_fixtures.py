"""Generates the binary fixtures (vsdx, xlsx, docx, compressed draw.io). Run: python3 make_fixtures.py"""
import base64, os, urllib.parse, zipfile, zlib
from openpyxl import Workbook
from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))

# --- draw.io: compressed page (raw deflate + URI encoding + base64) ---
src = open(os.path.join(HERE, 'drawio', 'beslissing-lus.source.xml'), encoding='utf-8').read()
co = zlib.compressobj(9, zlib.DEFLATED, -15)
packed = base64.b64encode(co.compress(urllib.parse.quote(src, safe='').encode()) + co.flush()).decode()
with open(os.path.join(HERE, 'drawio', 'beslissing-lus.drawio'), 'w', encoding='utf-8') as f:
    f.write(f'<mxfile host="app.diagrams.net"><diagram id="x" name="Lus">{packed}</diagram>'
            f'<diagram id="y" name="Tweede pagina">{packed}</diagram></mxfile>')

# --- Visio .vsdx ---
NS = 'http://schemas.microsoft.com/office/visio/2012/main'
MASTERS = {1: 'Start/End', 2: 'Process', 3: 'Decision', 4: 'Dynamic connector', 5: 'Swimlane'}

def shape(sid, master, text='', x=None, y=None, w=1.0, h=0.75, extra=''):
    cells = ''
    if x is not None:
        cells = f"<Cell N='PinX' V='{x}'/><Cell N='PinY' V='{y}'/><Cell N='Width' V='{w}'/><Cell N='Height' V='{h}'/>"
    body = f"<Text><cp IX='0'/>{text}</Text>" if text else ''
    return f"<Shape ID='{sid}' NameU='{MASTERS[master]}.{sid}' Type='Shape' Master='{master}'>{cells}{extra}{body}</Shape>"

def connector(sid, text=''):
    return shape(sid, 4, text, extra="<Cell N='BeginX' V='0'/><Cell N='EndX' V='1'/>")

def connect(cid, a, b):
    return (f"<Connect FromSheet='{cid}' FromCell='BeginX' FromPart='9' ToSheet='{a}' ToCell='PinX' ToPart='3'/>"
            f"<Connect FromSheet='{cid}' FromCell='EndX' FromPart='12' ToSheet='{b}' ToCell='PinX' ToPart='3'/>")

def vsdx(path, shapes, connects, pages=1):
    masters = ''.join(f"<Master ID='{i}' NameU='{n}' Name='{n}'><Rel r:id='rId{i}'/></Master>" for i, n in MASTERS.items())
    page = f"<?xml version='1.0' encoding='utf-8'?><PageContents xmlns='{NS}'><Shapes>{''.join(shapes)}</Shapes><Connects>{''.join(connects)}</Connects></PageContents>"
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr('[Content_Types].xml', "<?xml version='1.0'?><Types xmlns='http://schemas.openxmlformats.org/package/2006/content-types'/>")
        z.writestr('visio/document.xml', f"<?xml version='1.0'?><VisioDocument xmlns='{NS}'/>")
        z.writestr('visio/masters/masters.xml', f"<?xml version='1.0'?><Masters xmlns='{NS}' xmlns:r='http://schemas.openxmlformats.org/officeDocument/2006/relationships'>{masters}</Masters>")
        for n in range(1, pages + 1):
            z.writestr(f'visio/pages/page{n}.xml', page)

vsdx(os.path.join(HERE, 'vsdx', 'recht.vsdx'),
     [shape(1, 1, 'Start', 1, 5), shape(2, 2, 'Aanvraag ontvangen', 2.5, 5), shape(3, 2, 'Aanvraag beoordelen', 4, 5),
      shape(4, 2, 'Besluit versturen', 5.5, 5), shape(5, 1, 'Einde', 7, 5),
      connector(10), connector(11), connector(12), connector(13)],
     [connect(10, 1, 2), connect(11, 2, 3), connect(12, 3, 4), connect(13, 4, 5)])

vsdx(os.path.join(HERE, 'vsdx', 'beslissing-lus.vsdx'),
     [shape(1, 1, 'Start', 1, 5), shape(2, 2, 'Intake', 2.5, 5), shape(3, 2, 'Controle', 4, 5), shape(4, 3, 'Compleet?', 5.5, 5),
      shape(5, 2, 'Toekennen', 7, 5), shape(6, 2, 'Aanvullen', 5.5, 3), shape(7, 1, 'Einde', 8.5, 5),
      connector(10), connector(11), connector(12), connector(13, 'ja'), connector(14, 'nee'), connector(15), connector(16),
      connector(17, 'los')],
     [connect(10, 1, 2), connect(11, 2, 3), connect(12, 3, 4), connect(13, 4, 5), connect(14, 4, 6), connect(15, 6, 3), connect(16, 5, 7),
      "<Connect FromSheet='17' FromCell='BeginX' ToSheet='5'/>"],
     pages=2)

vsdx(os.path.join(HERE, 'vsdx', 'rollen.vsdx'),
     [shape(20, 5, 'Klantenservice', 5, 6, 10, 2), shape(21, 5, 'Backoffice', 5, 4, 10, 2), shape(22, 5, 'Teamleider', 5, 2, 10, 2),
      shape(1, 1, 'Start', 1, 6), shape(2, 2, 'Intake', 2.5, 6), shape(3, 2, 'Beoordelen', 4, 4), shape(4, 2, 'Goedkeuren', 5.5, 2),
      shape(5, 2, 'Versturen', 7, 6), shape(6, 1, 'Einde', 8.5, 6),
      connector(10), connector(11), connector(12), connector(13), connector(14)],
     [connect(10, 1, 2), connect(11, 2, 3), connect(12, 3, 4), connect(13, 4, 5), connect(14, 5, 6)])

with open(os.path.join(HERE, 'vsdx', 'kapot.vsdx'), 'wb') as f:
    f.write(b'PK\x03\x04 dit is geen geldig zipbestand')

# --- Excel ---
def xlsx(path, rows):
    wb = Workbook(); ws = wb.active
    for r in rows: ws.append(r)
    wb.save(path)

HEADER = ['id', 'stap', 'type', 'rol', 'systeem', 'bewerktijd_min', 'wachttijd_min', 'frequentie_per_jaar', 'foutpercentage', 'waardeklasse', 'volgende']
xlsx(os.path.join(HERE, 'table', 'recht.xlsx'), [['Stap', 'Rol'], ['Aanvraag ontvangen', 'Klantenservice'], ['Aanvraag beoordelen', 'Backoffice'], ['Besluit versturen', 'Klantenservice']])
xlsx(os.path.join(HERE, 'table', 'beslissing-lus.xlsx'), [HEADER,
    ['s', 'Start', 'start', None, None, None, None, None, None, None, 'intake'],
    ['intake', 'Intake', 'taak', 'Klantenservice', 'CRM', 10, 60, 1200, None, 'klantwaarde', 'controle'],
    ['controle', 'Controle', 'taak', 'Backoffice', 'ERP', 7.5, 120, 1200, 0.05, 'bedrijfsnoodzakelijk', 'd'],
    ['d', 'Compleet?', 'beslissing', None, None, None, None, None, None, None, 'toekennen:ja:0.8;aanvullen:nee:0.2'],
    ['toekennen', 'Toekennen', 'taak', 'Backoffice', 'ERP', 15, 30, 960, None, 'klantwaarde', 'e'],
    ['aanvullen', 'Aanvullen', 'taak', 'Klantenservice', None, 10, 1440, 240, None, 'geen waarde', 'controle'],
    ['e', 'Einde', 'einde', None, None, None, None, None, None, None, None]])
xlsx(os.path.join(HERE, 'table', 'rollen.xlsx'), [HEADER,
    ['i', 'Intake', '', 'Klantenservice', '', 5, 0, '', '', '', 'b'],
    ['b', 'Beoordelen', '', 'Backoffice', '', 20, 240, '', '', '', 'g'],
    ['g', 'Goedkeuren', '', 'Teamleider', '', 5, 480, '', '', '', 'v'],
    ['v', 'Versturen', '', 'Klantenservice', '', 5, 60, '', '', '', '']])

# --- Word ---
doc = Document()
doc.add_heading('Proces aanvraag', 1)
for line in ['De klantenservice ontvangt de aanvraag.', 'De backoffice beoordeelt de aanvraag.', 'De teamleider keurt goed.']:
    doc.add_paragraph(line)
doc.save(os.path.join(HERE, 'docx', 'proces.docx'))
print('fixtures written')
